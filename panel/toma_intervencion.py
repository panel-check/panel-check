"""
Toma de intervención para presentar ante INPI, una por marca (acta) del cliente.

Se arma desde la solapa Poderes de la ficha del cliente. El Word usa como plantilla
el modelo del estudio (recursos/toma-intervencion.docx): se reemplaza el texto de los
párrafos que cambian y se conserva todo el formato. El PDF se arma con el mismo texto.
Si son varias oposiciones, se genera un documento por cada una (el panel los junta en un ZIP).
"""

import io
import os
import re
import zipfile

import poderes

PLANTILLA = os.path.join(os.path.dirname(__file__), "recursos", "toma-intervencion.docx")

MATRICULA_DEFAULT = "2906"
EMAIL_DEFAULT = "smartiesconsultora@gmail.com"
TIPOS = {"D": "DENOMINATIVA", "M": "MIXTA", "F": "FIGURATIVA", "T": "TRIDIMENSIONAL"}


def matricula(apoderado: str) -> str:
    m = re.search(r"matr[ií]cula\s*(?:INPI\s*)?(?:N[°º.]?\s*)?(\d+)", apoderado or "", re.I)
    return m.group(1) if m else MATRICULA_DEFAULT


def tipo_texto(tipo) -> str:
    t = (tipo or "").strip()
    return TIPOS.get(t.upper(), t.upper())


def propuesta(marca: dict, cliente: dict, hoy) -> dict:
    """Valores sugeridos para una marca de la cartera; se pueden corregir en el panel."""
    return {"acta": str(marca.get("acta") or "").strip(), "marca": (marca.get("denominacion") or "").strip(),
            "clase": marca.get("clase") if marca.get("clase") is not None else "",
            "tipo": tipo_texto(marca.get("tipo")),
            "solicitante": ((marca.get("titular") or cliente.get("nombre") or "")).strip(),
            "fecha": hoy.isoformat()}


def _textos(d: dict, apoderado: str) -> dict:
    marca = (d.get("marca") or "").strip().upper()
    titulo = "TOMA DE INTERVENCION" + (f" “{marca}”" if marca else "")
    if str(d.get("clase") or "").strip():
        titulo += f" - CLASE {str(d['clase']).strip()}"
    if (d.get("tipo") or "").strip():
        titulo += f" - {d['tipo'].strip().upper()}"
    solicitante = (d.get("solicitante") or "").strip().rstrip(".")
    clase = f", Clase {str(d['clase']).strip()}" if str(d.get("clase") or "").strip() else ""
    return {
        "titulo": titulo,
        "fecha": f"Buenos Aires, {poderes.fecha_larga(d['fecha'])}",
        "cuerpo": (f"En mi carácter de Agente de la Propiedad Industrial, Matrícula INPI N° {matricula(apoderado)}, "
                   f"solicito se me tenga por presentada y por tomada intervención en el trámite correspondiente al "
                   f"Acta N° {str(d.get('acta') or '').strip()}{clase}, en representación del solicitante {solicitante}."),
        "domicilio": ("Asimismo, constituyo domicilio legal electrónico en el usuario registrado en el Portal de "
                      f"Trámites del INPI, asociado a la casilla {EMAIL_DEFAULT} donde se tendrán por válidas todas "
                      "las notificaciones del presente trámite."),
    }


def _poner(p, texto):
    """Reemplaza el texto del párrafo dejando el formato del primer fragmento."""
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(texto)


def generar_docx(d: dict, apoderado: str) -> bytes:
    from docx import Document
    t = _textos(d, apoderado)
    doc = Document(PLANTILLA)
    ps = doc.paragraphs
    for i, clave in ((0, "titulo"), (3, "fecha"), (9, "cuerpo"), (10, "domicilio")):
        _poner(ps[i], t[clave])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generar_pdf(d: dict, apoderado: str) -> bytes:
    from reportlab.lib.enums import TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    t = _textos(d, apoderado)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=3 * cm, rightMargin=3 * cm, topMargin=2.5 * cm,
                            bottomMargin=2.5 * cm, title="Toma de intervención", author="")
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=12, leading=16, alignment=TA_JUSTIFY)
    izq = ParagraphStyle("i", parent=base, alignment=0)
    der = ParagraphStyle("d", parent=base, alignment=TA_RIGHT)
    h = [Paragraph("<u>" + escape(t["titulo"]) + "</u>", base), Spacer(1, 1.2 * cm),
         Paragraph(escape(t["fecha"]), der), Spacer(1, 0.8 * cm),
         Paragraph("Señor", izq), Paragraph("Director de Marcas", izq), Spacer(1, 1.2 * cm),
         Paragraph(escape(t["cuerpo"]), base), Spacer(1, 0.4 * cm),
         Paragraph(escape(t["domicilio"]), base), Spacer(1, 1.2 * cm),
         Paragraph("Sin otro particular, saludo atentamente.", base)]
    doc.build(h)
    return buf.getvalue()


def nombre_archivo(d: dict, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚÑáéíóúñ]+", "_", f"{d.get('marca') or 'marca'}_acta_{d.get('acta') or ''}").strip("_")[:70]
    return f"Toma_de_intervencion_{base}.{ext}"


def generar_zip(items: list, apoderado: str, formato: str) -> bytes:
    gen = generar_docx if formato == "docx" else generar_pdf
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        usados = set()
        for d in items:
            n = nombre_archivo(d, formato)
            while n in usados:
                n = n.replace(f".{formato}", f"_2.{formato}")
            usados.add(n)
            z.writestr(n, gen(d, apoderado))
    return buf.getvalue()
