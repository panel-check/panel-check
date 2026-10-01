"""
Poder para el registro de marca, autocompletado con los datos del cliente.

Después de que el cliente completa el formulario (Clientes → Formularios),
desde su ficha se genera el poder ya lleno —en PDF o en Word— para mandárselo
y que solo tenga que firmarlo. Es el mismo texto del modelo del estudio
(persona física o jurídica); lo único que cambia es el párrafo del otorgante,
el lugar y fecha, y la aclaración/cargo de quien firma.

Los datos del apoderado (nombre, matrícula, domicilio) se editan en
Clientes → Ajustes y se guardan en vigilancia_config (clave poder_apoderado).
"""

import datetime as _dt
import io
import os
import re

LOGO = os.path.join(os.path.dirname(__file__), "recursos", "logo-poder.jpg")

APODERADO_DEFAULT = ("Pamela Daiana Guzzardi, agente de propiedad Industrial, Matrícula N° 2906, "
                     "domiciliada en la calle Oncativo 65, piso 1, departamento A, Ramos Mejía, República Argentina")

CUERPO = (
    "Poder amplio y suficiente para recabar de las autoridades nacionales que corresponda, sean éstas "
    "administrativas y / o judiciales, la obtención de Patentes de Invención y/o Marcas de Productos y "
    "Servicios y/o Modelos de Utilidad y/o Modelos Industriales, realizar todos los trámites conducentes "
    "para proteger los derechos de la mandante, a cuyo efecto, los faculta para efectuar ante dichas "
    "autoridades nacionales, administrativas y/o judiciales, todos los trámites necesarios a los objetos "
    "indicados; presentar solicitudes, alegatos, oposiciones, protestas, apelaciones y demandas; contestar "
    "vistas, exámenes, u observaciones a los trámites iniciados o a iniciarse; ampliar fundamentos de "
    "protestas y/o oposiciones y/o llamados de atención; desistir, arreglar y transigir; aceptar cesiones y "
    "transferencias; solicitar testimonios; recibir documentos y otros valores; justificar explotaciones; "
    "cobrar y percibir; firmar instrumentos públicos, privados y escrituras públicas; para renunciar o no a "
    "las gestiones judiciales; y hacer cuanto fuere necesario ante las autoridades nacionales, judiciales "
    "y/o administrativas, sea como demandante o como demandado, para proteger los intereses del "
    "demandante, con facultad de substituir este instrumento y revocar tal substitución."
)
CONSTANCIA = (
    "Se deja constancia que las facultades enumeradas en este mandato, son meramente ejemplificadas y no "
    "limitativas de otras facultades implícitas, por cuyo motivo, la mandataria queda autorizada para "
    "ejercer sin restricción alguna todos los derechos que acuerde la ley que se relacionen y/o emerjan y/o "
    "sean consecuencia del objeto del presente poder y ya sea ello en razón de la legislación vigente y/o de "
    "la legislación que en el futuro se promulgue."
)

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]
CABA = "Ciudad Autónoma de Buenos Aires"


def _dig(v):
    return re.sub(r"\D", "", str(v) if v is not None else "")


def fmt_dni(v):
    d = _dig(v)
    if not d:
        return ""
    return f"{int(d):,}".replace(",", ".")


def fmt_cuit(v):
    d = _dig(v)
    return f"{d[:2]}-{d[2:10]}-{d[10:]}" if len(d) == 11 else (v or "")


def fecha_larga(f) -> str:
    if isinstance(f, str):
        f = _dt.date.fromisoformat(f)
    return f"{f.day} de {MESES[f.month - 1]} de {f.year}"


def _es_caba(txt):
    t = (txt or "").lower()
    return "autónoma" in t or "autonoma" in t or t.strip() in ("caba", "capital federal")


def propuesta(cliente: dict, hoy=None) -> dict:
    """Valores sugeridos para el poder, a partir de la ficha del cliente.
    Todos se pueden corregir en el panel antes de generar."""
    hoy = hoy or _dt.date.today()
    c = cliente or {}
    tipo = c.get("tipo_persona") or ("juridica" if _dig(c.get("cuit")).startswith("3") else "fisica")
    nombre = (c.get("nombre") or "").strip()
    domicilio = str(c.get("domicilio") or "").strip().rstrip(".")
    provincia = (c.get("provincia") or "").strip()
    localidad = (c.get("localidad") or "").strip()
    if tipo == "juridica":
        partes = [f"{nombre}, CUIT {fmt_cuit(c.get('cuit'))}"]
        if domicilio:
            dom = domicilio
            if provincia and provincia.lower() not in domicilio.lower():
                dom += ", " + (CABA if _es_caba(provincia) else f"Provincia de {provincia}")
            partes.append(f"domiciliada en {dom}")
        otorgante = ", ".join(partes) + ", República Argentina."
        aclaracion = (c.get("firmante_nombre") or c.get("contacto") or "").strip()
        cargo = (c.get("firmante_cargo") or "").strip()
    else:
        lugar_dom = [domicilio] if domicilio else []
        if _es_caba(provincia):
            lugar_dom.append(CABA)
        else:
            if localidad and localidad.lower() not in domicilio.lower():
                lugar_dom.append(localidad)
            if provincia:
                lugar_dom.append(f"Provincia de {provincia}")
        partes = [nombre]
        if c.get("dni"):
            partes.append(f"DNI {fmt_dni(c.get('dni'))}")
        if lugar_dom:
            partes.append("domiciliado en " + ", ".join(lugar_dom))
        otorgante = ", ".join(partes) + ", República Argentina."
        aclaracion = nombre
        cargo = ""
    if _es_caba(provincia) or (tipo == "juridica" and _es_caba(domicilio)):
        lugar = CABA
    elif provincia:
        lugar = f"Provincia de {provincia}"
    else:
        lugar = CABA
    return {"tipo_persona": tipo, "otorgante": otorgante, "lugar": lugar, "fecha": hoy.isoformat(),
            "aclaracion": aclaracion, "cargo": cargo}


def _parrafo_inicial(d, apoderado):
    otorgante = d["otorgante"].strip().rstrip(".")
    return f"El abajo firmante {otorgante}. Por el presente otorgo a {apoderado.strip().rstrip('.')}. {CUERPO}"


def _lugar_fecha(d):
    return f"{d['lugar'].strip().rstrip(',')}, República Argentina, el {fecha_larga(d['fecha'])}." \
        if "república argentina" not in d["lugar"].lower() else f"{d['lugar'].strip()}, el {fecha_larga(d['fecha'])}."


def generar_pdf(d: dict, apoderado: str) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.6 * cm, rightMargin=2.2 * cm, topMargin=1.4 * cm,
                            bottomMargin=1.6 * cm, title="Poder", author="")
    cuerpo = ParagraphStyle("c", fontName="Times-Roman", fontSize=12, leading=14.6, alignment=TA_JUSTIFY)
    titulo = ParagraphStyle("t", parent=cuerpo, alignment=TA_CENTER)
    normal = ParagraphStyle("n", parent=cuerpo, alignment=0)
    h = []
    if os.path.exists(LOGO):
        img = Image(LOGO, width=3.6 * cm, height=3.6 * cm * 176 / 280)
        img.hAlign = "LEFT"
        h += [img, Spacer(1, 0.5 * cm)]
    else:
        h.append(Spacer(1, 2.5 * cm))
    h += [Paragraph("<u>Poder</u>", titulo), Spacer(1, 0.5 * cm),
          Paragraph(escape(_parrafo_inicial(d, apoderado)), cuerpo),
          Paragraph(escape(CONSTANCIA), cuerpo),
          Spacer(1, 1.9 * cm),
          Paragraph(escape(_lugar_fecha(d)), normal),
          Spacer(1, 1.8 * cm),
          Paragraph("FIRMA", normal), Spacer(1, 1.0 * cm),
          Paragraph("ACLARACIÓN: " + escape((d.get("aclaracion") or "").upper()), normal)]
    if d.get("tipo_persona") == "juridica" or d.get("cargo"):
        h += [Spacer(1, 0.7 * cm), Paragraph("CARGO: " + escape((d.get("cargo") or "").upper()), normal)]
    doc.build(h)
    return buf.getvalue()


def generar_docx(d: dict, apoderado: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2.2), Cm(1.4), Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.15

    def par(texto="", alin=None, antes=0, subrayado=False):
        p = doc.add_paragraph()
        if alin is not None:
            p.alignment = alin
        p.paragraph_format.space_before = Pt(antes)
        if texto:
            r = p.add_run(texto)
            r.underline = subrayado
        return p

    if os.path.exists(LOGO):
        doc.add_paragraph().add_run().add_picture(LOGO, width=Cm(3.6))
    par("Poder", WD_ALIGN_PARAGRAPH.CENTER, antes=12, subrayado=True)
    par(_parrafo_inicial(d, apoderado), WD_ALIGN_PARAGRAPH.JUSTIFY, antes=14)
    par(CONSTANCIA, WD_ALIGN_PARAGRAPH.JUSTIFY)
    par(_lugar_fecha(d), antes=54)
    par("FIRMA", antes=50)
    par("ACLARACIÓN: " + (d.get("aclaracion") or "").upper(), antes=28)
    if d.get("tipo_persona") == "juridica" or d.get("cargo"):
        par("CARGO: " + (d.get("cargo") or "").upper(), antes=20)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def nombre_archivo(d: dict, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚÑáéíóúñ]+", "_", (d.get("aclaracion") if d.get("tipo_persona") == "fisica" else "")
                  or d.get("otorgante", "").split(",")[0]).strip("_")[:60] or "cliente"
    return f"Poder_{base}.{ext}"
